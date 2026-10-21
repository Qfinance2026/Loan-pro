import os, io, csv, secrets
from datetime import date,datetime
from decimal import Decimal
from functools import wraps
from flask import Flask, request, session, redirect, url_for, render_template, flash, abort, send_file
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook, Workbook
from engine import D,q,make_schedule,state
from dotenv import load_dotenv
load_dotenv()
app=Flask(__name__)
if os.getenv('DATABASE_URL') and (not os.getenv('SECRET_KEY') or os.getenv('SECRET_KEY')=='local-development-only-change-me'):
    raise RuntimeError('Set strong SECRET_KEY in cloud environment')
app.config['SECRET_KEY']=os.getenv('SECRET_KEY','local-development-only-change-me')
url=os.getenv('DATABASE_URL','sqlite:///local_demo.sqlite3')
url=url.replace('postgres://','postgresql+psycopg://',1).replace('postgresql://','postgresql+psycopg://',1)
app.config['SQLALCHEMY_DATABASE_URI']=url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS']=False
app.config['SESSION_COOKIE_HTTPONLY']=True
app.config['SESSION_COOKIE_SAMESITE']='Lax'
app.config['SESSION_COOKIE_SECURE']=os.getenv('SESSION_COOKIE_SECURE','0')=='1'
app.config['PERMANENT_SESSION_LIFETIME']=__import__('datetime').timedelta(hours=8)
app.config['MAX_CONTENT_LENGTH']=5*1024*1024
if url.startswith('postgresql'):
    app.config['SQLALCHEMY_ENGINE_OPTIONS']={'pool_pre_ping':True,'pool_recycle':240,'pool_size':2,'max_overflow':0}
db=SQLAlchemy(app)
class User(db.Model):
    id=db.Column(db.Integer,primary_key=True);username=db.Column(db.String(80),unique=True,nullable=False)
    password=db.Column(db.String(255),nullable=False); role=db.Column(db.String(20),nullable=False,default='viewer');active=db.Column(db.Boolean,default=True)
class Borrower(db.Model):
    id=db.Column(db.Integer,primary_key=True);kind=db.Column(db.String(20),nullable=False);name=db.Column(db.String(200),nullable=False)
    registration=db.Column(db.String(50));phone=db.Column(db.String(60));address=db.Column(db.String(300))
class Loan(db.Model):
    id=db.Column(db.Integer,primary_key=True);borrower_id=db.Column(db.Integer,db.ForeignKey('borrower.id'),nullable=False)
    contract=db.Column(db.String(100),unique=True,nullable=False);amount=db.Column(db.Numeric(20,2),nullable=False)
    rate=db.Column(db.Numeric(9,4),nullable=False);penalty_rate=db.Column(db.Numeric(9,4),nullable=False,default=0)
    months=db.Column(db.Integer,nullable=False);start_date=db.Column(db.Date,nullable=False);method=db.Column(db.String(20),nullable=False)
    created_at=db.Column(db.DateTime,default=datetime.utcnow);borrower=db.relationship('Borrower')
class Due(db.Model):
    id=db.Column(db.Integer,primary_key=True);loan_id=db.Column(db.Integer,db.ForeignKey('loan.id'),nullable=False,index=True)
    version=db.Column(db.Integer,nullable=False,default=1);no=db.Column(db.Integer,nullable=False)
    due=db.Column(db.Date,nullable=False);principal=db.Column(db.Numeric(20,2),nullable=False)
    interest=db.Column(db.Numeric(20,2),nullable=False);total=db.Column(db.Numeric(20,2),nullable=False);balance=db.Column(db.Numeric(20,2),nullable=False)
class Payment(db.Model):
    id=db.Column(db.Integer,primary_key=True);loan_id=db.Column(db.Integer,db.ForeignKey('loan.id'),nullable=False,index=True)
    paid_at=db.Column(db.Date,nullable=False);amount=db.Column(db.Numeric(20,2),nullable=False)
    penalty_paid=db.Column(db.Numeric(20,2),nullable=False);interest_paid=db.Column(db.Numeric(20,2),nullable=False)
    principal_paid=db.Column(db.Numeric(20,2),nullable=False);note=db.Column(db.String(300))
    voided=db.Column(db.Boolean,default=False);created_by=db.Column(db.Integer,db.ForeignKey('user.id'));created_at=db.Column(db.DateTime,default=datetime.utcnow)
class Collateral(db.Model):
    id=db.Column(db.Integer,primary_key=True);loan_id=db.Column(db.Integer,db.ForeignKey('loan.id'),nullable=False)
    category=db.Column(db.String(100),nullable=False);description=db.Column(db.String(350));valuation=db.Column(db.Numeric(20,2),nullable=False)
class Audit(db.Model):
    id=db.Column(db.Integer,primary_key=True);user_id=db.Column(db.Integer);action=db.Column(db.String(60));entity=db.Column(db.String(60));entity_id=db.Column(db.Integer);details=db.Column(db.Text);created_at=db.Column(db.DateTime,default=datetime.utcnow)

def record(action,entity,eid,details=''):
    db.session.add(Audit(user_id=session.get('uid'),action=action,entity=entity,entity_id=eid,details=str(details)[:1000]))
def current_user():
    uid=session.get('uid');return db.session.get(User,uid) if uid else None
def protected(*roles):
    def deco(fn):
        @wraps(fn)
        def wrap(*args,**kwargs):
            u=current_user()
            if not u or not u.active:return redirect(url_for('login'))
            if roles and u.role not in roles:abort(403)
            return fn(*args,**kwargs)
        return wrap
    return deco
@app.before_request
def csrf():
    if 'csrf' not in session:session['csrf']=secrets.token_urlsafe(32)
    if request.method in ('POST','PUT','DELETE') and request.endpoint!='health':
        if not secrets.compare_digest(request.form.get('csrf',''),session['csrf']):abort(400,'CSRF validation failed')
@app.context_processor
def ctx():return dict(me=current_user(),csrf=session.get('csrf',''),today=date.today(),money=lambda x: f'{D(x):,.2f}',roles={'admin':'Админ','operator':'Ажилтан','viewer':'Харах'})
@app.errorhandler(403)
def forbidden(e):return 'Хандах эрх хүрэлцэхгүй',403
@app.route('/health')
def health():return {'status':'ok'}
@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='POST':
        u=User.query.filter_by(username=request.form.get('username','')).first()
        if u and u.active and check_password_hash(u.password,request.form.get('password','')):
            session.clear();session['uid']=u.id;session['csrf']=secrets.token_urlsafe(32);session.permanent=True
            return redirect(url_for('index'))
        flash('Нэвтрэх мэдээлэл буруу','error')
    return render_template('login.html')
@app.post('/logout')
def logout():session.clear();return redirect(url_for('login'))
def dues_of(loan):
    ver=db.session.query(func.max(Due.version)).filter_by(loan_id=loan.id).scalar()
    return Due.query.filter_by(loan_id=loan.id,version=ver).order_by(Due.due,Due.no).all()
def payments_of(loan):return Payment.query.filter_by(loan_id=loan.id).order_by(Payment.paid_at,Payment.id).all()
def state_of(loan,on=None):return state(loan,dues_of(loan),payments_of(loan),on or date.today())
@app.route('/')
@protected()
def index():
    loans=Loan.query.order_by(Loan.id.desc()).all()
    stats=[(l,state_of(l)) for l in loans]
    return render_template('index.html',stats=stats,total=sum((x['principal'] for l,x in stats),Decimal(0)), overdue=sum((x['overdue'] for l,x in stats),Decimal(0)),count=len(loans))
@app.route('/borrowers',methods=['GET','POST'])
@protected()
def borrowers():
    if request.method=='POST':
        if current_user().role not in ('admin','operator'):abort(403)
        name=request.form.get('name','').strip()
        if not name:flash('Нэр оруулна','error')
        else:
            b=Borrower(kind=request.form['kind'],name=name,registration=request.form.get('registration',''),phone=request.form.get('phone',''),address=request.form.get('address',''))
            db.session.add(b);db.session.flush();record('create','borrower',b.id,name);db.session.commit();return redirect(url_for('borrowers'))
    return render_template('borrowers.html',borrowers=Borrower.query.order_by(Borrower.id.desc()).all())
@app.route('/loans/new',methods=['GET','POST'])
@protected('admin','operator')
def new_loan():
    if request.method=='POST':
        try:
            amount=q(request.form['amount']);rate=D(request.form['rate']);penalty=D(request.form['penalty_rate']);months=int(request.form['months']);start=date.fromisoformat(request.form['start_date']);method=request.form['method']
            if penalty<0 or penalty>1000 or rate>1000 or amount>Decimal('999999999999999'):raise ValueError('Хүү эсвэл зээлийн дүн буруу')
            rows=make_schedule(amount,rate,start,months,method)
            if not db.session.get(Borrower,int(request.form['borrower_id'])):raise ValueError('Зээлдэгч олдсонгүй')
            loan=Loan(borrower_id=int(request.form['borrower_id']),contract=request.form['contract'].strip(),amount=amount,rate=rate,penalty_rate=penalty,months=months,start_date=start,method=method)
            if not loan.contract:raise ValueError('Гэрээний дугаар оруулна')
            db.session.add(loan);db.session.flush()
            for r in rows:db.session.add(Due(loan_id=loan.id,version=1,**r))
            record('create','loan',loan.id,loan.contract);db.session.commit()
            return redirect(url_for('loan_detail',loan_id=loan.id))
        except (ValueError,KeyError,IntegrityError) as e:
            db.session.rollback();flash('Зээл үүсгэсэнгүй: '+str(e)[:180],'error')
    return render_template('newloan.html',borrowers=Borrower.query.order_by(Borrower.name).all())
@app.route('/loans/<int:loan_id>')
@protected()
def loan_detail(loan_id):
    loan=db.get_or_404(Loan,loan_id);return render_template('loan.html',loan=loan,s=state_of(loan),dues=dues_of(loan),payments=payments_of(loan),collaterals=Collateral.query.filter_by(loan_id=loan_id).all())
@app.post('/loans/<int:loan_id>/payments')
@protected('admin','operator')
def pay(loan_id):
    loan=db.get_or_404(Loan,loan_id)
    try:
        paid_at=date.fromisoformat(request.form['date']);amount=q(request.form['amount']);mode=request.form.get('prepay','keep')
        if paid_at>date.today() or paid_at<loan.start_date or amount<=0:raise ValueError('Төлбөрийн дүн/огноо буруу')
        last=db.session.query(func.max(Payment.paid_at)).filter_by(loan_id=loan_id,voided=False).scalar()
        if last and paid_at<last:raise ValueError('Өмнөх гүйлгээнээс эрт өдөр бүртгэхгүй')
        s=state_of(loan,paid_at)
        if amount>s['total']:raise ValueError('Нийт өрөөс илүү төлөлт бүртгэхгүй')
        rem=amount;pen=min(rem,s['penalty']);rem-=pen;intr=min(rem,s['interest']);rem-=intr;principal=min(rem,s['principal'])
        p=Payment(loan_id=loan_id,paid_at=paid_at,amount=amount,penalty_paid=pen,interest_paid=intr,principal_paid=principal,note=request.form.get('note','')[:300],created_by=current_user().id)
        db.session.add(p);db.session.flush();record('payment','loan',loan_id,f'{paid_at} | {amount} | {pen}/{intr}/{principal}')
        if mode in ('reduce-payment','reduce-term') and principal>0 and s['principal']>principal:
            # Rescheduling based on actual balance, preserving prior dues as historical rows.
            prev=dues_of(loan); prior=[x for x in prev if x.due<=paid_at]; remaining=len(prev)-len(prior)
            if remaining:
                outstanding=q(s['principal']-principal);n=remaining
                if mode=='reduce-term':
                    old_payment=max((D(x.total) for x in prev if x.due>paid_at),default=D(0))
                    for candidate in range(1,remaining+1):
                        trial=make_schedule(outstanding,loan.rate,paid_at,candidate,loan.method)
                        if trial[0]['total']<=old_payment:n=candidate;break
                replacement=make_schedule(outstanding,loan.rate,paid_at,n,loan.method)
                ver=max(x.version for x in prev)+1
                for i,x in enumerate(prior):db.session.add(Due(loan_id=loan.id,version=ver,no=i+1,due=x.due,principal=x.principal,interest=x.interest,total=x.total,balance=x.balance))
                for i,r in enumerate(replacement):db.session.add(Due(loan_id=loan.id,version=ver,no=len(prior)+i+1,**{k:v for k,v in r.items() if k!='no'}))
                record('reschedule','loan',loan_id,mode)
        db.session.commit();flash('Төлөлт бүртгэгдлээ','ok')
    except (ValueError,KeyError) as e:db.session.rollback();flash('Алдаа: '+str(e),'error')
    return redirect(url_for('loan_detail',loan_id=loan_id))
@app.post('/loans/<int:loan_id>/collateral')
@protected('admin','operator')
def collateral(loan_id):
    db.get_or_404(Loan,loan_id)
    try:
        v=q(request.form['valuation'])
        if v<0:raise ValueError('Үнэлгээ сөрөг байж болохгүй')
        c=Collateral(loan_id=loan_id,category=request.form['category'][:100],description=request.form.get('description','')[:350],valuation=v)
        db.session.add(c);db.session.flush();record('collateral','loan',loan_id,c.category);db.session.commit()
    except (ValueError,KeyError) as e: db.session.rollback();flash(str(e),'error')
    return redirect(url_for('loan_detail',loan_id=loan_id))
@app.route('/users',methods=['GET','POST'])
@protected('admin')
def users():
    if request.method=='POST':
        username=request.form.get('username','').strip();password=request.form.get('password','');role=request.form.get('role','')
        if len(username)<3 or len(password)<12 or role not in ('admin','operator','viewer'):flash('Нэр 3+, нууц үг 12+ тэмдэгттэй байх ёстой','error')
        else:
            try:
                u=User(username=username,password=generate_password_hash(password),role=role);db.session.add(u);db.session.flush();record('create','user',u.id,role);db.session.commit()
            except IntegrityError:db.session.rollback();flash('Нэр давхцаж байна','error')
    return render_template('users.html',users=User.query.order_by(User.id).all())
@app.post('/users/<int:user_id>/toggle')
@protected('admin')
def toggle_user(user_id):
    u=db.get_or_404(User,user_id)
    if u.id==current_user().id:abort(400,'Өөрийн эрхийг хааж болохгүй')
    u.active=not u.active;record('toggle','user',u.id,str(u.active));db.session.commit();return redirect(url_for('users'))
@app.route('/reports')
@protected()
def reports():
    day=request.args.get('asof') or str(date.today())
    try:asof=date.fromisoformat(day)
    except ValueError:abort(400)
    rows=[(l,state_of(l,asof)) for l in Loan.query.order_by(Loan.id).all()]
    if request.args.get('export')=='xlsx':
        wb=Workbook();ws=wb.active;ws.title='Зээлийн багц';ws.append(['Гэрээ','Зээлдэгч','Үндсэн үлдэгдэл','Хүү','Нэмэгдүүлсэн хүү','Хэтэрсэн үндсэн','Хоног','Нийт өр'])
        for l,s in rows:ws.append([l.contract,l.borrower.name,*[float(s[k]) if k!='days' else s[k] for k in ('principal','interest','penalty','overdue','days','total')]])
        output=io.BytesIO();wb.save(output);output.seek(0)
        record('export','report',None,str(asof));db.session.commit()
        return send_file(output,as_attachment=True,download_name=f'loan_report_{asof}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    return render_template('reports.html',rows=rows,asof=asof)
@app.route('/import',methods=['GET','POST'])
@protected('admin','operator')
def import_excel():
    if request.method=='POST':
        f=request.files.get('file')
        if not f or not f.filename.lower().endswith('.xlsx'):flash('Зөвхөн .xlsx файл оруулна','error');return redirect(url_for('import_excel'))
        try:
            wb=load_workbook(io.BytesIO(f.read()),read_only=True,data_only=True);ws=wb.active
            expected=['kind','name','registration','phone','contract','amount','rate','penalty_rate','months','start_date','method']
            values=ws.iter_rows(values_only=True);head=next(values,None)
            if not head or list(head[:len(expected)])!=expected:raise ValueError('Excel загварын багана зөрсөн')
            imported=0
            for row in values:
                if not row or not row[0]:continue
                if imported>=500:raise ValueError('Нэг удаад 500 мөрөөс хэтрэхгүй')
                kind,name,reg,phone,contract,amount,rate,penrate,months,start,method=row[:11]
                if kind not in ('Иргэн','Байгууллага'):raise ValueError(f'{imported+2}-р мөрийн төрөл буруу')
                if isinstance(start,datetime):start=start.date()
                elif isinstance(start,date):pass
                else:start=date.fromisoformat(str(start)[:10])
                rows=make_schedule(amount,rate,start,int(months),str(method));penrate=D(penrate or 0)
                if penrate<0 or penrate>1000 or D(rate)>1000:raise ValueError('Хүү буруу')
                b=Borrower(kind=kind,name=str(name),registration=str(reg or ''),phone=str(phone or ''));db.session.add(b);db.session.flush()
                l=Loan(borrower_id=b.id,contract=str(contract),amount=q(amount),rate=D(rate),penalty_rate=penrate,months=int(months),start_date=start,method=str(method));db.session.add(l);db.session.flush()
                for r in rows:db.session.add(Due(loan_id=l.id,version=1,**r))
                record('import','loan',l.id,l.contract);imported+=1
            db.session.commit();flash(f'{imported} зээл импортоллоо','ok')
        except Exception as e:db.session.rollback();flash('Импорт амжилтгүй, бүх мөр буцаагдлаа: '+str(e)[:200],'error')
        return redirect(url_for('import_excel'))
    return render_template('import.html')
@app.get('/audit')
@protected('admin')
def audit():return render_template('audit.html',items=Audit.query.order_by(Audit.id.desc()).limit(150).all())

def initialize():
    with app.app_context():
        db.create_all()
        if User.query.count()==0:
            username=os.getenv('ADMIN_USERNAME','').strip();password=os.getenv('ADMIN_PASSWORD','')
            if len(username)>=3 and len(password)>=12:
                db.session.add(User(username=username,password=generate_password_hash(password),role='admin'));db.session.commit()
            elif os.getenv('DATABASE_URL'):
                app.logger.warning('INITIAL ADMIN NOT CREATED: configure ADMIN_USERNAME and ADMIN_PASSWORD (12+ chars).')
initialize()
if __name__=='__main__':app.run(debug=False)
